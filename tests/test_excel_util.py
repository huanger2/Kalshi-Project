import os
import tempfile
import unittest

import openpyxl
import pandas as pd

from ExcelUtil import write_workbook


class WriteWorkbookTests(unittest.TestCase):
    def test_writes_one_sheet_per_dataframe_in_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "out.xlsx")
            sheets = {
                "first": pd.DataFrame({"a": [1, 2]}),
                "second": pd.DataFrame({"b": [3]}),
            }
            result = write_workbook(sheets, out_path)

            self.assertEqual(result, out_path)
            wb = openpyxl.load_workbook(out_path)
            self.assertEqual(wb.sheetnames, ["first", "second"])
            self.assertEqual([c.value for c in wb["first"][1]], ["a"])
            self.assertEqual(wb["first"].max_row, 3)  # header + 2 rows
            self.assertEqual(wb["second"].max_row, 2)  # header + 1 row

    def test_empty_dataframe_still_writes_a_header_only_sheet(self):
        with tempfile.TemporaryDirectory() as tmp:
            out_path = os.path.join(tmp, "out.xlsx")
            write_workbook({"empty": pd.DataFrame(columns=["x", "y"])}, out_path)

            wb = openpyxl.load_workbook(out_path)
            self.assertEqual(wb["empty"].max_row, 1)
            self.assertEqual([c.value for c in wb["empty"][1]], ["x", "y"])


if __name__ == "__main__":
    unittest.main()
