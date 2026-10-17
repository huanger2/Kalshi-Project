import contextlib
import datetime
import io
import os
import tempfile
import unittest
from unittest import mock

import openpyxl
import pandas as pd

import DailyLog
import Tracker

OCT3 = datetime.date(2026, 10, 3)
OCT4 = datetime.date(2026, 10, 4)


def game_market(ticker, team, status="active", bid="0.60", ask="0.62"):
    return {
        "ticker": ticker,
        "event_ticker": ticker.rsplit("-", 1)[0],
        "title": f"{team} wins",
        "yes_sub_title": team,
        "status": status,
        "yes_bid_dollars": bid,
        "yes_ask_dollars": ask,
    }


TODAY_MARKETS = [
    game_market("KXNBAGAME-26OCT04UTADEN-DEN", "Denver"),
    game_market("KXNBAGAME-26OCT04UTADEN-UTA", "Utah", bid="0.38", ask="0.40"),
    game_market("KXNBAGAME-26OCT04GSWLAC-GSW", "Golden State", bid="0.58", ask="0.59"),
    game_market("KXNBAGAME-26OCT04GSWLAC-LAC", "Los Angeles C", bid="0.41", ask="0.42"),
]


def fake_llm(prob, conf):
    return lambda prompt: f"Probability: {prob}%\nConfidence: {conf}%"


LLMS = [("OpenAI", fake_llm(70, 80)), ("Claude", fake_llm(65, 60)), ("Gemini", fake_llm(55, 50))]


class DailyLogTestCase(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.folder = os.path.join(tmp.name, "daily_logs")
        self.path = DailyLog.path_for(OCT4, self.folder)
        # A real (but throwaway) file, not ":memory:", so a test can reopen
        # it afterward via Tracker.connect() to check what got recorded --
        # and, just as importantly, so run() never touches the real
        # predictions.db as a side effect of running the test suite.
        self.db_path = os.path.join(tmp.name, "predictions.db")

    def sheet(self, day):
        return DailyLog.load(DailyLog.path_for(day, self.folder))

    def run_day(self, today, markets=(), results=(), llms=LLMS):
        """Run DailyLog for `today` with Kalshi and CostExport mocked out; returns today's sheet."""
        with mock.patch.object(DailyLog, "get_todays_games", return_value=list(markets)), \
             mock.patch.object(DailyLog, "get_results", return_value=list(results)) as get_results, \
             mock.patch.object(DailyLog, "CostExport") as cost_export, \
             contextlib.redirect_stdout(io.StringIO()):
            self.counts = DailyLog.run(self.folder, today=today, llms=llms, db_path=self.db_path,
                                        leagues=["NBA"])
        self.get_results = get_results
        self.cost_export = cost_export
        return self.sheet(today)


class LogTodayTests(DailyLogTestCase):
    def test_one_row_per_game_with_odds_and_every_ai(self):
        df = self.run_day(OCT4, TODAY_MARKETS)

        self.assertEqual(len(df), 2)
        self.assertEqual(self.counts, (0, 2))
        row = df[df.event_ticker == "KXNBAGAME-26OCT04UTADEN"].iloc[0]
        self.assertEqual(row["game_date"], "2026-10-04")
        self.assertEqual(row["team"], "Denver")
        self.assertEqual(row["opponent"], "Utah")
        self.assertAlmostEqual(row["yes_bid"], 0.60)
        self.assertAlmostEqual(row["yes_ask"], 0.62)
        self.assertAlmostEqual(row["market_prob"], 0.61)
        self.assertAlmostEqual(row["openai_prob"], 0.70)
        self.assertAlmostEqual(row["openai_confidence"], 0.80)
        self.assertAlmostEqual(row["claude_prob"], 0.65)
        self.assertAlmostEqual(row["claude_confidence"], 0.60)
        self.assertAlmostEqual(row["gemini_prob"], 0.55)
        self.assertAlmostEqual(row["gemini_confidence"], 0.50)
        self.assertTrue(pd.isna(row["winner"]))
        self.assertTrue(pd.isna(row["team_won"]))

    def test_columns_are_in_a_stable_order(self):
        df = self.run_day(OCT4, TODAY_MARKETS)
        self.assertEqual(list(df.columns), DailyLog.columns(["OpenAI", "Claude", "Gemini"]))

    def test_rerunning_the_same_day_does_not_duplicate_games(self):
        self.run_day(OCT4, TODAY_MARKETS)
        df = self.run_day(OCT4, TODAY_MARKETS)
        self.assertEqual(len(df), 2)
        self.assertEqual(self.counts, (0, 0))

    def test_games_that_are_no_longer_open_are_skipped(self):
        markets = [
            game_market("KXNBAGAME-26OCT04UTADEN-DEN", "Denver", status="finalized"),
            game_market("KXNBAGAME-26OCT04UTADEN-UTA", "Utah", status="finalized"),
        ]
        df = self.run_day(OCT4, markets)
        self.assertEqual(len(df), 0)
        self.assertEqual(self.counts, (0, 0))

    def test_a_failing_ai_leaves_its_cells_blank(self):
        def broken(prompt):
            raise RuntimeError("API down")

        llms = [("OpenAI", broken), ("Claude", fake_llm(65, 60)), ("Gemini", fake_llm(55, 50))]
        row = self.run_day(OCT4, TODAY_MARKETS[:2], llms=llms).iloc[0]
        self.assertTrue(pd.isna(row["openai_prob"]))
        self.assertAlmostEqual(row["claude_prob"], 0.65)

    def test_game_is_not_logged_when_every_ai_fails(self):
        def broken(prompt):
            raise RuntimeError("API down")

        llms = [("OpenAI", broken), ("Claude", broken), ("Gemini", broken)]
        df = self.run_day(OCT4, TODAY_MARKETS[:2], llms=llms)
        self.assertEqual(len(df), 0)  # so tomorrow's run doesn't treat it as done


class FillResultsTests(DailyLogTestCase):
    RESULTS = [{"event_ticker": "KXNBAGAME-26OCT04UTADEN", "winner": "Utah", "loser": "Denver"}]

    def test_next_day_fills_in_yesterdays_results(self):
        self.run_day(OCT4, TODAY_MARKETS)
        self.run_day(datetime.date(2026, 10, 5), results=self.RESULTS)
        df = self.sheet(OCT4)  # results go in the file for the day the game was played

        self.get_results.assert_called_once_with(OCT4, "NBA")
        row = df[df.event_ticker == "KXNBAGAME-26OCT04UTADEN"].iloc[0]
        self.assertEqual(row["winner"], "Utah")
        self.assertEqual(row["team_won"], 0)  # the row is for Denver, who lost
        self.assertFalse(pd.isna(row["settled_at"]))
        self.assertEqual(self.counts, (1, 0))

    def test_team_won_is_1_when_the_logged_team_won(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])
        results = [{"event_ticker": "KXNBAGAME-26OCT04UTADEN", "winner": "Denver", "loser": "Utah"}]
        self.run_day(datetime.date(2026, 10, 5), results=results)
        self.assertEqual(self.sheet(OCT4).iloc[0]["team_won"], 1)

    def test_unfinished_games_stay_pending(self):
        self.run_day(OCT4, TODAY_MARKETS)
        self.run_day(datetime.date(2026, 10, 5), results=self.RESULTS)
        df = self.sheet(OCT4)
        row = df[df.event_ticker == "KXNBAGAME-26OCT04GSWLAC"].iloc[0]
        self.assertTrue(pd.isna(row["winner"]))

    def test_a_skipped_day_is_caught_up_later(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])
        self.run_day(datetime.date(2026, 10, 5), results=[])  # result not in yet
        self.run_day(datetime.date(2026, 10, 7), results=self.RESULTS)
        self.get_results.assert_called_once_with(OCT4, "NBA")
        self.assertEqual(self.sheet(OCT4).iloc[0]["winner"], "Utah")

    def test_todays_own_games_are_not_settled_yet(self):
        self.run_day(OCT4, TODAY_MARKETS)
        self.run_day(OCT4, TODAY_MARKETS, results=self.RESULTS)
        self.get_results.assert_not_called()

    def test_settled_games_are_not_looked_up_again(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])
        self.run_day(datetime.date(2026, 10, 5), results=self.RESULTS)
        self.run_day(datetime.date(2026, 10, 6), results=self.RESULTS)
        self.get_results.assert_not_called()


class DayFileTests(DailyLogTestCase):
    def test_todays_games_go_in_a_file_named_for_today_in_the_folder(self):
        self.run_day(OCT4, TODAY_MARKETS)
        self.assertEqual(os.listdir(self.folder), ["2026-10-04.xlsx"])

    def test_each_day_gets_its_own_file(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])
        oct5 = [game_market("KXNBAGAME-26OCT05LALSAC-LAL", "Los Angeles L"),
                game_market("KXNBAGAME-26OCT05LALSAC-SAC", "Sacramento")]
        self.run_day(datetime.date(2026, 10, 5), oct5)

        self.assertEqual(sorted(os.listdir(self.folder)), ["2026-10-04.xlsx", "2026-10-05.xlsx"])
        self.assertEqual(list(self.sheet(OCT4).team), ["Denver"])
        self.assertEqual(list(self.sheet(datetime.date(2026, 10, 5)).team), ["Los Angeles L"])

    def test_no_file_is_created_on_a_day_without_games(self):
        self.run_day(OCT4, [])
        self.assertEqual(os.listdir(self.folder), [])

    def test_other_files_in_the_folder_are_ignored(self):
        os.makedirs(self.folder)
        for name in ("~$2026-10-03.xlsx", "notes.xlsx", "2026-10-03.csv"):
            open(os.path.join(self.folder, name), "w").close()
        self.assertEqual(DailyLog.earlier_days(self.folder, OCT4), [])


class SaveTests(unittest.TestCase):
    def test_file_open_in_excel_gives_a_clear_message(self):
        with mock.patch.object(openpyxl.Workbook, "save", side_effect=PermissionError):
            with self.assertRaises(SystemExit) as ctx:
                DailyLog.save(pd.DataFrame(), "daily_log.xlsx")
        self.assertIn("close it in Excel", str(ctx.exception))


class LayoutTests(DailyLogTestCase):
    def setUp(self):
        super().setUp()
        self.run_day(OCT4, TODAY_MARKETS)
        self.wb = openpyxl.load_workbook(self.path)
        self.ws = self.wb["Daily log"]

    def row_values(self, row):
        return [c.value for c in self.ws[row]]

    def test_row_1_groups_columns_by_source(self):
        groups = [v for v in self.row_values(1) if v]
        self.assertEqual(groups, ["Game", "Kalshi odds", "OpenAI", "Claude", "Gemini", "Result", "Details"])
        merged = {str(r) for r in self.ws.merged_cells.ranges}
        self.assertIn("H1:J1", merged)  # OpenAI spans its three columns (win %, confidence, model)

    def test_row_2_has_readable_headers(self):
        headers = self.row_values(2)
        self.assertEqual(headers[:8], [
            "Game date", "League", "Team", "Opponent", "Kalshi bid", "Kalshi ask",
            "Kalshi implied win %", "OpenAI win %",
        ])
        self.assertIn("Claude confidence", headers)
        self.assertIn("Team won? (1 = yes)", headers)

    def test_probabilities_are_formatted_as_percentages(self):
        headers = self.row_values(2)
        for header in ("Kalshi implied win %", "Claude win %", "Gemini confidence"):
            cell = self.ws.cell(3, headers.index(header) + 1)
            self.assertEqual(cell.number_format, DailyLog.PERCENT, header)
            self.assertLessEqual(cell.value, 1)  # stored as a fraction, shown as %

    def test_each_source_has_its_own_color(self):
        headers = self.row_values(2)
        color = lambda h: self.ws.cell(2, headers.index(h) + 1).fill.fgColor.rgb
        self.assertNotEqual(color("Kalshi bid"), color("OpenAI win %"))
        self.assertNotEqual(color("OpenAI win %"), color("Claude win %"))
        self.assertEqual(color("Claude win %"), color("Claude confidence"))

    def test_headers_stay_visible_when_scrolling(self):
        self.assertEqual(self.ws.freeze_panes, "E3")

    def test_column_guide_explains_every_column_and_its_source(self):
        guide = self.wb["Column guide"]
        rows = [[c.value for c in r] for r in guide.iter_rows(min_row=2)]
        self.assertEqual([r[0] for r in rows], self.row_values(2))
        sources = {r[0]: r[1] for r in rows}
        self.assertEqual(sources["Kalshi ask"], "Kalshi")
        self.assertEqual(sources["Gemini win %"], "Gemini")
        self.assertTrue(all(r[2] for r in rows))

    def test_formatted_sheet_reads_back_to_the_same_data(self):
        df = DailyLog.load(self.path)
        self.assertEqual(list(df.columns), DailyLog.columns(["OpenAI", "Claude", "Gemini"]))
        row = df[df.event_ticker == "KXNBAGAME-26OCT04UTADEN"].iloc[0]
        self.assertEqual(row["game_date"], "2026-10-04")
        self.assertAlmostEqual(row["claude_prob"], 0.65)


class OldLayoutTests(DailyLogTestCase):
    def test_single_header_sheet_from_before_is_still_read(self):
        old = pd.DataFrame([{
            "game_date": "2026-10-04", "event_ticker": "KXNBAGAME-26OCT04UTADEN",
            "ticker": "KXNBAGAME-26OCT04UTADEN-DEN", "team": "Denver", "opponent": "Utah",
            "market_prob": 0.61, "openai_prob": 0.7, "winner": None,
        }])
        os.makedirs(self.folder)
        old.to_excel(self.path, index=False)

        results = [{"event_ticker": "KXNBAGAME-26OCT04UTADEN", "winner": "Denver", "loser": "Utah"}]
        self.run_day(datetime.date(2026, 10, 5), results=results)
        df = self.sheet(OCT4)
        self.assertEqual(df.iloc[0]["team_won"], 1)
        self.assertAlmostEqual(df.iloc[0]["openai_prob"], 0.7)
        self.assertEqual(openpyxl.load_workbook(self.path)["Daily log"]["A2"].value, "Game date")


class DatabaseIntegrationTests(DailyLogTestCase):
    """
    predictions.db is the table the webapp queries; DailyLog.run() must keep
    it in sync with the spreadsheet as games are logged and settled.
    """

    def query_db(self, sql):
        # Windows can't clean up the test's temp directory while a
        # connection to a file in it is still open, so this must close
        # before returning.
        conn = Tracker.connect(self.db_path)
        try:
            return conn.execute(sql).fetchall()
        finally:
            conn.close()

    def test_logging_a_game_records_one_row_per_ai_in_the_database(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])  # one game: Denver vs Utah

        rows = self.query_db(
            "SELECT llm, team, opponent, predicted_prob, confidence, market_prob, game_date, outcome "
            "FROM predictions ORDER BY llm"
        )
        self.assertEqual(rows, [
            ("Claude", "Denver", "Utah", 0.65, 0.60, 0.61, "2026-10-04", None),
            ("Gemini", "Denver", "Utah", 0.55, 0.50, 0.61, "2026-10-04", None),
            ("OpenAI", "Denver", "Utah", 0.70, 0.80, 0.61, "2026-10-04", None),
        ])

    def test_settling_a_game_updates_the_database_row_too(self):
        self.run_day(OCT4, TODAY_MARKETS[:1])
        results = [{"event_ticker": "KXNBAGAME-26OCT04UTADEN", "winner": "Utah", "loser": "Denver"}]
        self.run_day(datetime.date(2026, 10, 5), results=results)

        outcomes = {r[0] for r in self.query_db("SELECT outcome FROM predictions")}
        self.assertEqual(outcomes, {0})  # the row is for Denver, who lost

    def test_a_database_failure_while_logging_does_not_lose_the_spreadsheet_row(self):
        with mock.patch.object(Tracker, "record_predictions", side_effect=RuntimeError("db locked")):
            df = self.run_day(OCT4, TODAY_MARKETS[:1])
        self.assertEqual(len(df), 1)
        self.assertAlmostEqual(df.iloc[0]["claude_prob"], 0.65)

    def test_a_database_failure_while_settling_does_not_lose_the_spreadsheet_result(self):
        self.run_day(OCT4, TODAY_MARKETS[:1])
        results = [{"event_ticker": "KXNBAGAME-26OCT04UTADEN", "winner": "Utah", "loser": "Denver"}]
        with mock.patch.object(Tracker, "settle", side_effect=RuntimeError("db locked")):
            self.run_day(datetime.date(2026, 10, 5), results=results)
        self.assertEqual(self.sheet(OCT4).iloc[0]["winner"], "Utah")


class ModelColumnTests(DailyLogTestCase):
    def test_todays_rows_records_the_configured_model_for_each_successful_ai(self):
        with mock.patch.object(DailyLog, "get_model", side_effect=lambda p: f"{p}-test-model"):
            df = self.run_day(OCT4, TODAY_MARKETS[:2])
        row = df.iloc[0]
        self.assertEqual(row["openai_model"], "openai-test-model")
        self.assertEqual(row["claude_model"], "claude-test-model")
        self.assertEqual(row["gemini_model"], "gemini-test-model")

    def test_a_failing_ai_has_no_model_recorded(self):
        def broken(prompt):
            raise RuntimeError("API down")

        llms = [("OpenAI", broken), ("Claude", fake_llm(65, 60)), ("Gemini", fake_llm(55, 50))]
        with mock.patch.object(DailyLog, "get_model", side_effect=lambda p: f"{p}-test-model"):
            df = self.run_day(OCT4, TODAY_MARKETS[:2], llms=llms)
        self.assertTrue(pd.isna(df.iloc[0]["openai_model"]))
        self.assertEqual(df.iloc[0]["claude_model"], "claude-test-model")

    def test_backfill_records_the_model_too(self):
        def broken(prompt):
            raise RuntimeError("API down")

        llms = [("OpenAI", broken), ("Claude", fake_llm(65, 60)), ("Gemini", fake_llm(55, 50))]
        self.run_day(OCT4, TODAY_MARKETS[:1], llms=llms)

        with mock.patch.object(DailyLog, "get_model", return_value="new-openai-model"):
            DailyLog.backfill_missing_predictions(
                OCT4, self.folder, llms=[("OpenAI", fake_llm(80, 90))],
                fetch_market=lambda t: next(m for m in TODAY_MARKETS if m["ticker"] == t),
            )
        self.assertEqual(self.sheet(OCT4).iloc[0]["openai_model"], "new-openai-model")


class BackfillLayoutRegressionTests(DailyLogTestCase):
    """
    Regression guard: an earlier version passed the (possibly --only=
    restricted) queried-AI list straight to save() for column *layout* too,
    which shunted an untouched AI's columns into a generic "Other" group at
    the end of the sheet instead of leaving them in place.
    """

    def test_backfilling_just_one_ai_does_not_misplace_the_others_columns(self):
        self.run_day(OCT4, TODAY_MARKETS[:1])  # all 3 AIs log normally

        DailyLog.backfill_missing_predictions(
            OCT4, self.folder, llms=[("Gemini", fake_llm(99, 1))],
            fetch_market=lambda t: next(m for m in TODAY_MARKETS if m["ticker"] == t), force=True,
        )

        wb = openpyxl.load_workbook(self.path)
        groups = [v for v in [c.value for c in wb["Daily log"][1]] if v]
        self.assertEqual(groups, ["Game", "Kalshi odds", "OpenAI", "Claude", "Gemini", "Result", "Details"])
        self.assertNotIn("Other", groups)


class BackfillTests(DailyLogTestCase):
    def market_lookup(self):
        """fetch_market stand-in: looks a ticker up in TODAY_MARKETS instead of hitting Kalshi."""
        return lambda ticker: next(m for m in TODAY_MARKETS if m["ticker"] == ticker)

    def test_fills_in_only_the_missing_llms_cells(self):
        def broken(prompt):
            raise RuntimeError("API down")

        llms = [("OpenAI", broken), ("Claude", fake_llm(65, 60)), ("Gemini", fake_llm(55, 50))]
        two_games = [TODAY_MARKETS[0], TODAY_MARKETS[2]]  # Denver and Golden State -- different events
        self.run_day(OCT4, two_games, llms=llms)  # OpenAI blank for both games

        n = DailyLog.backfill_missing_predictions(
            OCT4, self.folder, llms=[("OpenAI", fake_llm(80, 90))], fetch_market=self.market_lookup(),
        )

        self.assertEqual(n, 2)
        df = self.sheet(OCT4)
        self.assertTrue((df["openai_prob"] == 0.80).all())
        self.assertTrue((df["openai_confidence"] == 0.90).all())
        self.assertAlmostEqual(df.iloc[0]["claude_prob"], 0.65)  # untouched

    def test_does_nothing_when_nothing_is_missing(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])  # every AI succeeds
        n = DailyLog.backfill_missing_predictions(OCT4, self.folder)
        self.assertEqual(n, 0)

    def test_a_missing_day_file_is_a_no_op(self):
        n = DailyLog.backfill_missing_predictions(OCT4, self.folder)
        self.assertEqual(n, 0)

    def test_a_still_failing_llm_leaves_the_cell_blank_and_is_reported(self):
        def broken(prompt):
            raise RuntimeError("API down")

        self.run_day(OCT4, TODAY_MARKETS[:1], llms=[
            ("OpenAI", broken), ("Claude", fake_llm(65, 60)), ("Gemini", fake_llm(55, 50)),
        ])

        with contextlib.redirect_stdout(io.StringIO()):
            n = DailyLog.backfill_missing_predictions(
                OCT4, self.folder, llms=[("OpenAI", broken)], fetch_market=self.market_lookup(),
            )

        self.assertEqual(n, 0)
        self.assertTrue(pd.isna(self.sheet(OCT4).iloc[0]["openai_prob"]))

    def test_rows_that_are_already_filled_in_are_not_refetched(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])  # every AI already succeeded
        fetch = mock.Mock()
        n = DailyLog.backfill_missing_predictions(OCT4, self.folder, llms=[("OpenAI", fake_llm(1, 1))], fetch_market=fetch)
        self.assertEqual(n, 0)
        fetch.assert_not_called()

    def test_force_replaces_every_row_even_when_already_filled(self):
        two_games = [TODAY_MARKETS[0], TODAY_MARKETS[2]]  # different events
        self.run_day(OCT4, two_games)  # every AI succeeds with the original LLMS fixture

        n = DailyLog.backfill_missing_predictions(
            OCT4, self.folder, llms=[("Gemini", fake_llm(99, 88))],
            fetch_market=self.market_lookup(), force=True,
        )

        self.assertEqual(n, 2)
        df = self.sheet(OCT4)
        self.assertTrue((df["gemini_prob"] == 0.99).all())
        self.assertTrue((df["gemini_confidence"] == 0.88).all())
        # Untouched: force only applies to the LLMs passed in.
        self.assertAlmostEqual(df.iloc[0]["claude_prob"], 0.65)

    def test_without_force_an_already_filled_llm_is_left_alone(self):
        self.run_day(OCT4, TODAY_MARKETS[:2])
        n = DailyLog.backfill_missing_predictions(
            OCT4, self.folder, llms=[("Gemini", fake_llm(99, 88))], fetch_market=self.market_lookup(),
        )
        self.assertEqual(n, 0)
        self.assertAlmostEqual(self.sheet(OCT4).iloc[0]["gemini_prob"], 0.55)  # original LLMS value


NFL_MARKETS = [
    game_market("KXNFLGAME-26OCT04DALPHI-PHI", "Philadelphia"),
    game_market("KXNFLGAME-26OCT04DALPHI-DAL", "Dallas", bid="0.38", ask="0.40"),
]
NHL_MARKETS = [
    game_market("KXNHLGAME-26OCT04TBCOL-TB", "Tampa Bay"),
    game_market("KXNHLGAME-26OCT04TBCOL-COL", "Colorado", bid="0.38", ask="0.40"),
]


class MultiLeagueTests(DailyLogTestCase):
    BY_LEAGUE = {"NBA": TODAY_MARKETS[:2], "NFL": NFL_MARKETS, "NHL": NHL_MARKETS, "MLB": []}

    def run_leagues(self, today, leagues, results_by_league=None):
        results_by_league = results_by_league or {}
        with mock.patch.object(DailyLog, "get_todays_games",
                               side_effect=lambda league, day: list(self.BY_LEAGUE[league]) if day == OCT4 else []), \
             mock.patch.object(DailyLog, "get_results",
                               side_effect=lambda day, league: list(results_by_league.get(league, []))) as gr, \
             mock.patch.object(DailyLog, "CostExport"), \
             contextlib.redirect_stdout(io.StringIO()):
            DailyLog.run(self.folder, today=today, llms=LLMS, db_path=self.db_path, leagues=leagues)
        self.get_results = gr
        return self.sheet(today)

    def test_logs_every_requested_league_into_one_days_file(self):
        df = self.run_leagues(OCT4, ["NBA", "NFL", "NHL", "MLB"])
        self.assertEqual(sorted(df["league"]), ["NBA", "NFL", "NHL"])
        nfl = df[df.league == "NFL"].iloc[0]
        self.assertEqual((nfl["team"], nfl["opponent"]), ("Philadelphia", "Dallas"))

    def test_only_the_requested_leagues_are_queried(self):
        df = self.run_leagues(OCT4, ["NHL"])
        self.assertEqual(list(df["league"]), ["NHL"])

    def test_defaults_to_the_leagues_enabled_in_leagues_json(self):
        with mock.patch.object(DailyLog.Leagues, "enabled_leagues", return_value=["NFL"]):
            df = self.run_leagues(OCT4, None)
        self.assertEqual(list(df["league"]), ["NFL"])

    def test_results_are_fetched_per_league(self):
        self.run_leagues(OCT4, ["NBA", "NFL"])
        self.run_leagues(datetime.date(2026, 10, 5), ["NBA", "NFL"])
        self.assertEqual(sorted(c.args[1] for c in self.get_results.call_args_list), ["NBA", "NFL"])

    def test_an_nfl_tie_is_recorded_as_no_team_winning(self):
        self.run_leagues(OCT4, ["NFL"])
        tie = [{"event_ticker": "KXNFLGAME-26OCT04DALPHI", "winner": None, "loser": None, "tie": True}]
        self.run_leagues(datetime.date(2026, 10, 5), ["NFL"], {"NFL": tie})

        row = self.sheet(OCT4).iloc[0]
        self.assertEqual(row["winner"], "Tie")
        self.assertEqual(row["team_won"], 0)
        conn = Tracker.connect(self.db_path)
        self.addCleanup(conn.close)
        self.assertEqual(set(conn.execute("SELECT outcome, tied FROM predictions")), {(0, 1)})

    def test_league_is_recorded_in_predictions_db(self):
        self.run_leagues(OCT4, ["NFL", "NHL"])
        conn = Tracker.connect(self.db_path)
        self.addCleanup(conn.close)
        self.assertEqual({r[0] for r in conn.execute("SELECT DISTINCT league FROM predictions")}, {"NFL", "NHL"})

    def test_a_file_from_before_the_league_column_reads_as_nba(self):
        old = pd.DataFrame([{
            "game_date": "2026-10-04", "event_ticker": "KXNBAGAME-26OCT04UTADEN",
            "ticker": "KXNBAGAME-26OCT04UTADEN-DEN", "team": "Denver", "opponent": "Utah",
            "market_prob": 0.61, "openai_prob": 0.7, "winner": None,
        }])
        os.makedirs(self.folder)
        old.to_excel(self.path, index=False)
        self.assertEqual(DailyLog.load(self.path).iloc[0]["league"], "NBA")


class CostExportIntegrationTests(DailyLogTestCase):
    def test_cost_spreadsheet_is_refreshed_after_every_run(self):
        self.run_day(OCT4, TODAY_MARKETS)
        self.cost_export.export_to_excel.assert_called_once()

    def test_cost_spreadsheet_is_refreshed_even_with_no_new_games(self):
        self.run_day(OCT4, [])
        self.cost_export.export_to_excel.assert_called_once()

    def test_a_locked_cost_spreadsheet_does_not_fail_the_whole_run(self):
        with mock.patch.object(DailyLog, "get_todays_games", return_value=list(TODAY_MARKETS)), \
             mock.patch.object(DailyLog, "get_results", return_value=[]), \
             mock.patch.object(DailyLog, "CostExport") as cost_export, \
             contextlib.redirect_stdout(io.StringIO()):
            cost_export.export_to_excel.side_effect = PermissionError("locked")
            counts = DailyLog.run(self.folder, today=OCT4, llms=LLMS, db_path=self.db_path, leagues=["NBA"])
        # Games were still logged despite the cost-export step failing.
        self.assertEqual(counts, (0, 2))


if __name__ == "__main__":
    unittest.main()
