"""
DailyLog.py

Run once at the start of the day to collect data before live betting:

  1. Fill in the results of earlier games that are still pending (yesterday's,
     plus any day that was missed), in the file for the day they were played.
  2. Log today's games in every league enabled in leagues.json (NBA, NFL,
     NHL, MLB): Kalshi's current odds and each AI's predicted probability
     and confidence.

Each day gets its own spreadsheet in daily_logs/, named by date
(daily_logs/2026-10-04.xlsx), with one row per game across every league. The
row is for one team (the YES side of the market that was asked about);
team_won is 1 if that team won, 0 if it lost or tied (NFL only -- no team
won), blank until the game is settled. Re-running on the same day skips games that
are already logged.

Each sheet groups its columns by source (Game, Kalshi odds, each AI, Result)
under a colored header row, and a 'Column guide' sheet explains every column.

    py DailyLog.py [folder]
    py DailyLog.py backfill [YYYY-MM-DD] [folder] [--force] [--only=OpenAI,Gemini]

A game is only retried automatically if *every* AI failed for it; if one AI
fails while the others succeed, the game is still logged, with that AI's
cells left permanently blank. `backfill` fixes an already-logged day: for
each AI with a blank cell, it re-fetches that market and re-queries just
that AI, leaving everything else untouched.

Add `--force` to re-query every row instead of just the blank cells -- for
replacing predictions logged under an AI's old model with fresh ones after
upgrading it in ai_settings.json. Combine with `--only=` to limit that to
specific AIs, since --force re-spends on every row for the ones it covers.
"""

import contextlib
import datetime
import itertools
import os
import sqlite3
import sys
from dataclasses import dataclass
from typing import Callable, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import Betting
import CostExport
import DataDir
import Leagues
import Tracker
from AISettings import get_model
from Kalshi import get_market, get_prices, get_results, get_todays_games, implied_yes_prob
from Probability import build_prompt, default_llms, parse_response, query_llms

DEFAULT_DIR = DataDir.data_path("daily_logs")
DEFAULT_LLM_NAMES = ["OpenAI", "Claude", "Gemini"]


# ---------------------------------------------------------------------------
# Columns: internal key, spreadsheet header, the source group it sits under,
# and the explanation shown on the "Column guide" sheet.
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Column:
    key: str
    header: str
    group: str
    source: str
    meaning: str
    number_format: Optional[str] = None


PERCENT = "0.0%"
GROUP_COLORS = {"Game": "D9D9D9", "Kalshi odds": "C6E0B4", "Result": "FFE699", "Details": "EDEDED"}
AI_COLORS = ["BDD7EE", "F8CBAD", "D9C3E9", "B4E5E0", "FFD3E0"]


def column_spec(llm_names: list[str]) -> list[Column]:
    spec = [
        Column("game_date", "Game date", "Game", "Kalshi",
               "Date the game is scheduled (US time)."),
        Column("league", "League", "Game", "Kalshi",
               "NBA, NFL, NHL or MLB."),
        Column("team", "Team", "Game", "Kalshi",
               "The team this row is about. Every win % on the row is the chance THIS team wins."),
        Column("opponent", "Opponent", "Game", "Kalshi",
               "The other team in the game."),
        Column("yes_bid", "Kalshi bid", "Kalshi odds", "Kalshi",
               "Highest price buyers were offering for a 'Team wins' contract when the row was "
               "logged, as a % of its $1 payout.", PERCENT),
        Column("yes_ask", "Kalshi ask", "Kalshi odds", "Kalshi",
               "Lowest price a 'Team wins' contract could be bought for when the row was logged. "
               "This is what betting on Team would have cost.", PERCENT),
        Column("market_prob", "Kalshi implied win %", "Kalshi odds", "Kalshi",
               "Midpoint of the bid and ask: the market's estimate of the chance Team wins.", PERCENT),
    ]
    for name in llm_names:
        key = name.lower()
        spec += [
            Column(f"{key}_prob", f"{name} win %", name, name,
                   f"{name}'s estimate of the chance Team wins (model set in ai_settings.json).", PERCENT),
            Column(f"{key}_confidence", f"{name} confidence", name, name,
                   f"How sure {name} says it is of its estimate: 0% = pure guess, 100% = certain.", PERCENT),
            Column(f"{key}_model", f"{name} model", name, name,
                   f"Which {name} model produced this row's prediction. Lets you tell an old "
                   f"model's predictions apart from a newer one's after an ai_settings.json upgrade."),
        ]
    spec += [
        Column("winner", "Winner", "Result", "Kalshi",
               "The team that won, or 'Tie' (NFL only). Filled in by the first run after the "
               "game finishes."),
        Column("team_won", "Team won? (1 = yes)", "Result", "Kalshi",
               "1 if Team won, 0 if Team lost or tied (NFL only -- a tie means no team won, and "
               "Winner says 'Tie'), blank until the game is settled."),
        Column("settled_at", "Result recorded at", "Result", "This script",
               "When the result was filled in."),
        Column("logged_at", "Logged at", "Details", "This script",
               "When the Kalshi odds and AI answers on this row were captured."),
        Column("event_ticker", "Kalshi game ID", "Details", "Kalshi",
               "Kalshi's ID for the game (covers both teams)."),
        Column("ticker", "Kalshi market ID", "Details", "Kalshi",
               "Kalshi's ID for the 'Team wins' market whose odds are on this row."),
    ]
    return spec


def columns(llm_names: list[str]) -> list[str]:
    return [c.key for c in column_spec(llm_names)]


def _now() -> str:
    return datetime.datetime.now().isoformat(timespec="seconds")


# ---------------------------------------------------------------------------
# Reading and writing the workbook
# ---------------------------------------------------------------------------

def _header_to_key(header: str) -> str:
    """Map a spreadsheet header back to its key (any AI name works)."""
    fixed = {c.header: c.key for c in column_spec([])}
    if header in fixed:
        return fixed[header]
    if header.endswith(" win %"):
        return header[: -len(" win %")].lower() + "_prob"
    if header.endswith(" confidence"):
        return header[: -len(" confidence")].lower() + "_confidence"
    if header.endswith(" model"):
        return header[: -len(" model")].lower() + "_model"
    return header


def _is_numeric(key: str) -> bool:
    return key in ("yes_bid", "yes_ask", "team_won") or key.endswith(("_prob", "_confidence"))


def load(path: str) -> pd.DataFrame:
    """The log so far, or an empty one. Dates are kept as ISO strings."""
    if not os.path.exists(path):
        return pd.DataFrame(columns=columns(DEFAULT_LLM_NAMES))

    # Everything as text first: columns that are still all blank would
    # otherwise load as floats and reject team names later.
    if "game_date" in pd.read_excel(path, nrows=0).columns:
        df = pd.read_excel(path, dtype=object)  # older single-header layout
    else:
        df = pd.read_excel(path, header=1, dtype=object)  # row 1 is the source groups
        df = df.rename(columns=_header_to_key)

    for key in df.columns:
        if _is_numeric(key):
            df[key] = pd.to_numeric(df[key])

    # Files from before the League column existed (all NBA): the ticker says.
    if "league" not in df.columns:
        df.insert(min(1, len(df.columns)), "league", None)
    missing = df["league"].isna()
    df.loc[missing, "league"] = [
        Leagues.league_for_ticker(t) or Leagues.DEFAULT_LEAGUE for t in df.loc[missing, "ticker"]
    ]
    return df


def _cell_value(value):
    if pd.isna(value):
        return None
    return value.item() if hasattr(value, "item") else value  # numpy -> plain Python


def save(df: pd.DataFrame, path: str, llm_names: list[str] = DEFAULT_LLM_NAMES) -> None:
    """
    Write the log as a formatted workbook: row 1 groups the columns by source,
    row 2 names each column, and a 'Column guide' sheet explains them all.
    """
    spec = column_spec(llm_names)
    known = {c.key for c in spec}
    spec += [Column(k, k, "Other", "", "") for k in df.columns if k not in known]
    df = df.reindex(columns=[c.key for c in spec])

    colors = dict(GROUP_COLORS, Other="FFFFFF")
    colors.update({name: AI_COLORS[i % len(AI_COLORS)] for i, name in enumerate(llm_names)})

    bold = Font(bold=True)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    divider = Border(left=Side(style="medium", color="808080"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Daily log"

    # Row 1: one merged, colored cell per source group.
    col = 1
    group_starts = set()
    for group, members in itertools.groupby(spec, key=lambda c: c.group):
        width = len(list(members))
        group_starts.add(col)
        ws.cell(1, col, group)
        if width > 1:
            ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + width - 1)
        for c in range(col, col + width):
            ws.cell(1, c).fill = PatternFill("solid", fgColor=colors[group])
        ws.cell(1, col).font = bold
        ws.cell(1, col).alignment = centered
        col += width

    # Row 2: column names, colored to match their group.
    for i, column in enumerate(spec, 1):
        cell = ws.cell(2, i, column.header)
        cell.font = bold
        cell.alignment = centered
        cell.fill = PatternFill("solid", fgColor=colors[column.group])
        ws.column_dimensions[get_column_letter(i)].width = max(11, min(len(column.header) + 3, 22))
    ws.row_dimensions[2].height = 32

    # Data from row 3.
    for r, record in enumerate(df.itertuples(index=False), 3):
        for i, (value, column) in enumerate(zip(record, spec), 1):
            cell = ws.cell(r, i, _cell_value(value))
            if column.number_format:
                cell.number_format = column.number_format

    # A divider line down the left edge of each group.
    for i in group_starts - {1}:
        for r in range(1, ws.max_row + 1):
            ws.cell(r, i).border = divider

    ws.freeze_panes = "E3"  # keep the headers and Game columns visible
    ws.auto_filter.ref = f"A2:{get_column_letter(len(spec))}{max(ws.max_row, 2)}"

    guide = wb.create_sheet("Column guide")
    for i, (header, width) in enumerate([("Column", 24), ("Source", 14), ("What it means", 95)], 1):
        cell = guide.cell(1, i, header)
        cell.font = bold
        guide.column_dimensions[get_column_letter(i)].width = width
    for r, column in enumerate((c for c in spec if c.group != "Other"), 2):
        guide.cell(r, 1, column.header).fill = PatternFill("solid", fgColor=colors[column.group])
        guide.cell(r, 2, column.source)
        guide.cell(r, 3, column.meaning).alignment = Alignment(wrap_text=True, vertical="top")
    guide.freeze_panes = "A2"

    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"Can't write {path}; close it in Excel and run again.") from None


def fill_results(
    df: pd.DataFrame,
    today: datetime.date,
    fetch_results: Optional[Callable[[datetime.date, str], list[dict]]] = None,
    db_conn: Optional[sqlite3.Connection] = None,
) -> int:
    """
    Fill winner / team_won for pending games before `today`. Returns rows settled.

    If `db_conn` is given, the matching predictions.db rows are settled too
    (see Tracker.settle()) so the webapp's accuracy numbers stay in sync with
    the spreadsheet. A failure there is printed, not raised -- the result
    still belongs in the spreadsheet even if the database write fails.
    """
    fetch_results = fetch_results or get_results
    pending = df["winner"].isna() & (df["game_date"] < today.isoformat())
    settled = 0
    for (day, league), group in df[pending].groupby(["game_date", "league"]):
        results = {r["event_ticker"]: r for r in fetch_results(datetime.date.fromisoformat(day), league)}
        for idx in group.index:
            result = results.get(df.at[idx, "event_ticker"])
            if result is None:
                continue
            tie = bool(result.get("tie"))
            if tie:
                winner, team_won = "Tie", 0
            else:
                winner, team_won = result["winner"], int(df.at[idx, "team"] == result["winner"])
            df.at[idx, "winner"] = winner
            df.at[idx, "team_won"] = team_won
            df.at[idx, "settled_at"] = _now()
            settled += 1
            if db_conn is not None:
                try:
                    Tracker.settle(db_conn, df.at[idx, "ticker"], bool(team_won), tie=tie)
                except Exception as e:
                    print(f"could not settle {df.at[idx, 'ticker']} in predictions.db: {e}")
    return settled


def todays_rows(
    today: datetime.date,
    already_logged: set[str],
    llms: list[tuple[str, Callable[[str], str]]],
    db_conn: Optional[sqlite3.Connection] = None,
    leagues: Optional[list[str]] = None,
) -> list[dict]:
    """
    One row per game today, in each of `leagues` (default: every league
    enabled in leagues.json), that isn't logged yet and is still open for
    trading.

    If `db_conn` is given, each game's predictions are also stored in
    predictions.db (see Tracker.record_predictions()) -- the table the
    webapp queries. A failure there is printed, not raised: the spreadsheet
    row is still the thing that matters if the database write fails.
    """
    games: dict[str, list[dict]] = {}
    for league in leagues or Leagues.enabled_leagues():
        for market in get_todays_games(league, today):
            games.setdefault(market["event_ticker"], []).append(market)

    bet_rules = None
    if db_conn is not None:
        try:
            bet_rules = Betting.load_config()
        except Exception as e:
            print(f"no bets today: could not read the betting rules: {e}")
    bets = []

    rows = []
    for event, markets in games.items():
        if event in already_logged:
            continue
        market = markets[0]
        if market.get("status") != "active":
            print(f"skipped {event}: market is {market.get('status')}, not open")
            continue

        predictions = query_llms(market, llms, verbose=False)
        if not predictions:
            # Leave it unlogged so the next run tries again.
            print(f"skipped {event}: every AI query failed")
            continue

        prices = get_prices(market)
        market_prob = implied_yes_prob(market)
        opponent = next((m.get("yes_sub_title") for m in markets[1:]), None)
        row = {
            "game_date": today.isoformat(),
            "league": Leagues.league_for_ticker(market["ticker"]),
            "event_ticker": event,
            "ticker": market["ticker"],
            "team": market.get("yes_sub_title"),
            "opponent": opponent,
            "logged_at": _now(),
            "yes_bid": prices["yes_bid"],
            "yes_ask": prices["yes_ask"],
            "market_prob": market_prob,
        }
        for p in predictions:
            key = p.llm_name.lower()
            row[f"{key}_prob"] = p.predicted_prob
            row[f"{key}_confidence"] = p.confidence
            row[f"{key}_model"] = get_model(key)
        rows.append(row)
        print(f"logged {event}: {row['team']} vs {row['opponent']} ({row['league']})")

        if db_conn is not None:
            try:
                Tracker.record_predictions(
                    db_conn, market, market_prob, predictions,
                    opponent=opponent, yes_bid=prices["yes_bid"], yes_ask=prices["yes_ask"],
                    game_date=today.isoformat(),
                )
            except Exception as e:
                print(f"could not record {event} in predictions.db: {e}")
            # Paper bets (betting.json) -- best-effort too: a betting problem
            # mustn't cost the game its spreadsheet row or its predictions.
            if bet_rules is not None:
                try:
                    bet = Betting.candidate(db_conn, market, predictions, prices, today.isoformat(),
                                            opponent=opponent, config=bet_rules)
                    if bet:
                        bets.append(bet)
                except Exception as e:
                    print(f"could not consider a bet on {event}: {e}")

    # Placed once every game is in, so the daily cap goes to the biggest edges.
    if bets:
        try:
            Betting.place(db_conn, bets, bet_rules)
        except Exception as e:
            print(f"could not place today's bets: {e}")
    return rows


def backfill_missing_predictions(
    day: datetime.date,
    folder: str = DEFAULT_DIR,
    llms: Optional[list[tuple[str, Callable[[str], str]]]] = None,
    fetch_market=get_market,
    force: bool = False,
) -> int:
    """
    Fill in AI prediction cells for an already-logged day.

    By default only fills still-blank cells: todays_rows() only retries a
    game if *every* AI failed for it, so a game where one AI failed and the
    others succeeded gets logged anyway, leaving that AI's cells permanently
    blank.

    With force=True, every row is re-queried for the given `llms` regardless
    of whether a cell is already filled -- for re-running a day's games
    through a newly-upgraded model to replace predictions logged under the
    old one. Pass a specific `llms` list to limit this to just the AI(s) you
    changed; it's deliberately not the default, since it re-spends on every
    row instead of only the ones that actually failed.

    Either way, this re-fetches the market for each cell touched and
    re-queries just that AI, leaving every other cell untouched.

    Returns the number of cells filled in.
    """
    llms = llms or default_llms()
    path = path_for(day, folder)
    df = load(path)
    if df.empty:
        return 0

    filled = 0
    for llm_name, ask_fn in llms:
        key = llm_name.lower()
        prob_col, conf_col, model_col = f"{key}_prob", f"{key}_confidence", f"{key}_model"
        if prob_col not in df.columns:
            continue

        targets = df.index if force else df.index[df[prob_col].isna()]
        for idx in targets:
            ticker = df.at[idx, "ticker"]
            try:
                prompt = build_prompt(fetch_market(ticker))
                prob, conf = parse_response(ask_fn(prompt), llm_name)
            except Exception as e:
                print(f"[{llm_name}] could not backfill {ticker}: {e}")
                continue
            df.at[idx, prob_col] = prob
            df.at[idx, conf_col] = conf
            df.at[idx, model_col] = get_model(key)
            filled += 1
            print(f"backfilled {llm_name} for {df.at[idx, 'team']} ({ticker}): "
                  f"{prob:.0%} probability, {conf:.0%} confidence")

    if filled:
        # Always the full canonical layout, not just the (possibly
        # --only=-restricted) subset of AIs that were re-queried this call --
        # otherwise an untouched AI's columns get shunted into a generic
        # "Other" group at the end of the sheet instead of staying in place.
        save(df, path, DEFAULT_LLM_NAMES)
    return filled


def path_for(day: datetime.date, folder: str = DEFAULT_DIR) -> str:
    """The spreadsheet for one day's games, e.g. daily_logs/2026-10-04.xlsx."""
    return os.path.join(folder, f"{day.isoformat()}.xlsx")


def earlier_days(folder: str, today: datetime.date) -> list[datetime.date]:
    """Dates of the day files in `folder` from before `today`, oldest first."""
    days = []
    for name in os.listdir(folder):
        stem, ext = os.path.splitext(name)
        try:
            day = datetime.date.fromisoformat(stem)  # also skips Excel's ~$ lock files
        except ValueError:
            continue
        if ext == ".xlsx" and day < today:
            days.append(day)
    return sorted(days)


def run(
    folder: str = DEFAULT_DIR,
    today: Optional[datetime.date] = None,
    llms: Optional[list[tuple[str, Callable[[str], str]]]] = None,
    db_path: str = Tracker.DEFAULT_DB,
    leagues: Optional[list[str]] = None,
) -> tuple[int, int]:
    """
    Fill in results in earlier days' files, then write today's games in
    each of `leagues` (default: every league enabled in leagues.json) to
    today's file. Returns (games settled, games logged).

    Every game is also recorded in predictions.db (db_path) as it's logged
    or settled -- the single table the webapp/ package queries, so it never
    falls behind the spreadsheet. Pass db_path=":memory:" (or any throwaway
    path) to skip persisting that -- e.g. in tests that only care about the
    spreadsheet.
    """
    today = today or datetime.date.today()
    llms = llms or default_llms()
    llm_names = [name for name, _ in llms]
    os.makedirs(folder, exist_ok=True)
    db_conn = Tracker.connect(db_path)

    settled = 0
    for day in earlier_days(folder, today):
        path = path_for(day, folder)
        df = load(path)
        n = fill_results(df, today, db_conn=db_conn)
        if n:
            save(df, path, llm_names)
            settled += n

    # Backstop: settle predictions.db rows whose day file isn't in `folder`
    # -- e.g. games logged on another machine before this one took over the
    # daily run -- straight from Kalshi, so they never sit pending forever.
    # Best-effort, like the spreadsheet-side settling above.
    try:
        settled += len(Tracker.settle_from_kalshi(db_conn, fetch_market=get_market, before=today.isoformat()))
    except Exception as e:
        print(f"could not settle remaining games from Kalshi: {e}")

    path = path_for(today, folder)
    df = load(path)
    new_rows = todays_rows(today, set(df["event_ticker"].dropna()), llms, db_conn=db_conn, leagues=leagues)
    if new_rows:
        new = pd.DataFrame(new_rows)
        df = new if df.empty else pd.concat([df, new], ignore_index=True)
        save(df, path, llm_names)

    print(f"settled {settled} earlier game(s); logged {len(new_rows)} game(s) to {path}")

    # Every AI call today was already logged to predictions.db as it happened
    # (see CostTracker.py); refresh the running api_costs.xlsx from it so the
    # spreadsheet never falls behind without needing a separate manual step.
    # Best-effort: a locked/open spreadsheet shouldn't fail the whole daily run.
    try:
        cost_path = CostExport.export_to_excel()
        print(f"updated {cost_path}")
    except Exception as e:
        print(f"could not update cost spreadsheet: {e}")

    db_conn.close()  # Windows can't reopen/export predictions.db while this is open
    return settled, len(new_rows)


@contextlib.contextmanager
def _prevent_sleep():
    """
    Hold the system awake for this run. schedule_task.ps1's -WakeToRun wakes
    a sleeping machine to start this script, but that alone doesn't stop
    Windows' own idle timer from putting it back to sleep a few minutes in
    -- observed in practice: the 9am run got cut off mid-way, logged nothing
    for the day, and left predictions.db a day stale. ES_SYSTEM_REQUIRED
    tells Windows this process needs the system awake; cleared again on
    exit so it doesn't keep the machine up after the run finishes. No-op on
    non-Windows.
    """
    if sys.platform != "win32":
        yield
        return
    import ctypes
    ES_CONTINUOUS = 0x80000000
    ES_SYSTEM_REQUIRED = 0x00000001
    ctypes.windll.kernel32.SetThreadExecutionState(ES_CONTINUOUS | ES_SYSTEM_REQUIRED)
    try:
        yield
    finally:
        ctypes.windll.kernel32.SetThreadExecutionState(ES_CONTINUOUS)


if __name__ == "__main__":
    with _prevent_sleep():
        if len(sys.argv) > 1 and sys.argv[1] == "backfill":
            args = sys.argv[2:]
            force = "--force" in args
            only = next((a.split("=", 1)[1] for a in args if a.startswith("--only=")), None)
            positional = [a for a in args if not a.startswith("--")]

            day = datetime.date.fromisoformat(positional[0]) if positional else datetime.date.today()
            folder = positional[1] if len(positional) > 1 else DEFAULT_DIR

            llms = default_llms()
            if only:
                names = {n.strip().lower() for n in only.split(",")}
                llms = [(name, fn) for name, fn in llms if name.lower() in names]

            n = backfill_missing_predictions(day, folder, llms=llms, force=force)
            print(f"backfilled {n} cell(s)" if n else "nothing to backfill")
        else:
            run(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_DIR)
